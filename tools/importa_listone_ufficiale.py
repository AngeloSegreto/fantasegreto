#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, re, unicodedata
from pathlib import Path
from openpyxl import load_workbook

REQUIRED = ["Id","R","RM","Nome","Squadra","Qt.A","Qt.I","FVM"]
ROLES = {"P","D","C","A"}

def sha256_file(p: Path) -> str:
    h=hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024), b""):
            h.update(chunk)
    return h.hexdigest()

def norm_header(x):
    return str(x).strip() if x is not None else ""

def parse_name(raw):
    raw=unicodedata.normalize("NFC", str(raw or "")).strip()
    starred=raw.endswith("*")
    clean=raw[:-1].rstrip() if starred else raw
    return raw, clean, starred

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--output", required=True)
    ap.add_argument("--receipt")
    args=ap.parse_args()
    src=Path(args.xlsx).resolve()
    out=Path(args.output).resolve()
    wb=load_workbook(src, read_only=True, data_only=True)
    if "Tutti" not in wb.sheetnames:
        raise SystemExit("BLOCK: foglio 'Tutti' assente")
    ws=wb["Tutti"]
    headers=[norm_header(c.value) for c in next(ws.iter_rows())]
    idx={h:i for i,h in enumerate(headers) if h}
    miss=[h for h in REQUIRED if h not in idx]
    if miss:
        raise SystemExit("BLOCK: header mancanti: "+", ".join(miss))
    seen=set(); rows=[]; errors=[]
    for rno,row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals=list(row)
        if all(v in (None,"") for v in vals):
            continue
        def g(h):
            i=idx[h]
            return vals[i] if i < len(vals) else None
        oid=str(g("Id")).strip() if g("Id") is not None else ""
        role=str(g("R") or "").strip().upper()
        rm=str(g("RM") or "").strip()
        raw,name,starred=parse_name(g("Nome"))
        team=str(g("Squadra") or "").strip()
        if not oid or oid in seen: errors.append(f"riga {rno}: Id assente/duplicato {oid!r}")
        seen.add(oid)
        if role not in ROLES: errors.append(f"riga {rno}: ruolo {role!r}")
        if not rm: errors.append(f"riga {rno}: RM mancante")
        try: qa=float(g("Qt.A"))
        except: errors.append(f"riga {rno}: Qt.A non numerico"); qa=None
        try: qi=float(g("Qt.I"))
        except: errors.append(f"riga {rno}: Qt.I non numerico"); qi=None
        try: fvm=float(g("FVM"))
        except: errors.append(f"riga {rno}: FVM non numerico"); fvm=None
        rows.append({
            "id_ufficiale":oid,"nome_raw":raw,"nome":name,"starred":starred,
            "ruolo":role,"ruolo_mantra":rm,"squadra":team,
            "quotazione":qa,"quotazione_iniziale":qi,"fvm":fvm
        })
    if errors:
        raise SystemExit("BLOCK parser:\n" + "\n".join(errors[:50]))
    payload={
      "schema":"FS_LISTONE_CANONICAL_V3_9_6_CANDIDATE",
      "source_file":src.name,
      "source_sha256":sha256_file(src),
      "count":len(rows),
      "players":rows
    }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    if args.receipt:
        rp=Path(args.receipt)
        rp.parent.mkdir(parents=True, exist_ok=True)
        rp.write_text(json.dumps({
          "schema":"FS_SOURCE_RECEIPT_V1",
          "source":"Fantacalcio.it official XLSX",
          "filename":src.name,
          "sha256":payload["source_sha256"],
          "bytes":src.stat().st_size,
          "records":len(rows)
        }, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(f"PASS candidate={out} records={len(rows)} sha256={payload['source_sha256']}")

if __name__=="__main__":
    main()
